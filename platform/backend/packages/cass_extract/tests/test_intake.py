"""The CASS intake template: profile, generator and reader.

The property that matters most is the binding. A column exists so that a person
can type "Building value" and the platform can write ``BuildingTIV``, and every
test here that looks pedantic is protecting that: a column bound to a field OED
does not define, or a type validated differently from the field it lands in,
would be a mapping that reads correctly and imports wrongly.

The second property is that a blank means something specific and is never zero.
"""

from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from cass_extract import intake, profile, template
from cass_extract.profile import ProfileError, Sheet, WhenBlank
from cass_oed.schema import SCHEMA, DataType, FileKind


def risk(**overrides):
    row = {
        "Policy ID": "A1",
        "Risk reference": "1",
        "Country": "ID",
        "Latitude": "-6.2088",
        "Longitude": "106.8456",
        "Perils covered": "QEQ",
        "Currency": "USD",
        "Building value": "1000000.00",
    }
    row.update(overrides)
    return row


def policy(**overrides):
    row = {
        "Policy ID": "A1",
        "Policy reference": "P1",
        "Currency": "USD",
        "Perils covered": "QEQ",
    }
    row.update(overrides)
    return row


def read(risks=(), policies=(), contracts=(), scope=()) -> intake.IntakeRead:
    payload = template.workbook(
        risks=list(risks),
        policies=list(policies),
        contracts=list(contracts),
        scope=list(scope),
    )
    return intake.read_workbook(io.BytesIO(payload))


def contract(**overrides):
    row = {
        "Contract number": "2",
        "Layer": "1",
        "Contract type": "CXL",
        "Inuring priority": "1",
        "Attachment per event": "5000000",
        "Limit per event": "20000000",
    }
    row.update(overrides)
    return row


# -- the binding ---------------------------------------------------------------------

def test_the_profile_is_internally_consistent():
    """Run as a test so a broken mapping fails a build, not a user's upload."""
    assert list(profile.validate()) == []


def test_every_column_lands_in_a_field_oed_actually_defines():
    for item in profile.COLUMNS:
        if item.oed_field is None:
            continue
        known = {spec.name for spec in SCHEMA[item.oed_kind]}
        assert item.oed_field in known, f"{item.name} -> {item.oed_field}"


def test_each_sheet_lands_in_its_own_oed_file():
    assert {
        sheet: {item.oed_kind for item in profile.columns_for(sheet) if item.oed_kind}
        for sheet in Sheet
    } == {
        Sheet.RISK: {FileKind.LOCATION},
        Sheet.POLICY: {FileKind.ACCOUNT},
        Sheet.CONTRACT: {FileKind.REINS_INFO},
        Sheet.SCOPE: {FileKind.REINS_SCOPE},
    }


def test_a_column_is_read_as_the_type_of_the_field_it_lands_in():
    """Otherwise a column could validate as text and be written as money."""
    for item in profile.COLUMNS:
        spec = item.spec
        if spec is not None:
            assert item.reads_as is spec.dtype


def test_every_oed_required_field_is_asked_for_or_derived():
    """A template that could not produce a valid OED file would be a trap."""
    for kind, fields in SCHEMA.items():
        bound = {item.oed_field for item in profile.COLUMNS if item.oed_kind is kind}
        for spec in fields:
            if not spec.required:
                continue
            assert spec.name in bound or spec.name in profile.DERIVED_FIELDS


def test_the_binding_is_published_as_data():
    """A loading API follows this rather than reimplementing the mapping."""
    binding = profile.oed_binding(Sheet.RISK)
    assert binding["Building value"] == "BuildingTIV"
    assert binding["Latitude"] == "Latitude"
    assert "Risk name" not in binding  # CASS-only, no OED destination


def test_the_fields_the_template_does_not_cover_are_named():
    """Silence about a gap reads as coverage."""
    uncovered = {spec.name for spec in profile.unmapped_oed_fields(FileKind.LOCATION)}
    assert uncovered == {"IsTenant", "LocGroup"}
    assert profile.unmapped_oed_fields(FileKind.ACCOUNT) == ()


def test_a_column_with_no_oed_field_must_say_what_it_feeds():
    with pytest.raises(ProfileError, match="no OED field and no purpose"):
        profile.Column(
            name="Orphan",
            sheet=Sheet.RISK,
            oed_field=None,
            oed_kind=None,
            required=False,
            help_text="",
            when_blank=WhenBlank.ABSENT,
        )


def test_a_required_column_may_not_declare_a_fallback():
    """A required column with a fallback is an optional column."""
    with pytest.raises(ProfileError, match="required but declares a behaviour"):
        profile.Column(
            name="Country",
            sheet=Sheet.RISK,
            oed_field="CountryCode",
            oed_kind=FileKind.LOCATION,
            required=True,
            help_text="",
            when_blank=WhenBlank.ASSUMED,
        )


def test_every_cass_only_column_states_its_own_type():
    for item in profile.COLUMNS:
        if item.oed_field is None:
            assert item.dtype is not None
            assert item.purpose


def test_no_column_is_classified():
    """A Klapton Re portfolio holds nothing a Klapton Re colleague may not see.

    An earlier draft graded columns by sensitivity and role-gated the address.
    That was an inference from a rule about counterparty *names*, and it had a
    cost: a modeller who cannot read an address cannot check a coordinate
    against it, which is the whole of the geocoding review.
    """
    assert not hasattr(profile.Column, "is_confidential")
    assert "confidential" not in profile.COLUMNS[0].as_dict()


# -- the generated template ------------------------------------------------------------

def test_the_template_carries_the_guidance_and_four_sheets():
    book = load_workbook(io.BytesIO(template.workbook()))
    assert book.sheetnames == [
        profile.GUIDE_SHEET,
        profile.RISK_SHEET,
        profile.POLICY_SHEET,
        profile.CONTRACT_SHEET,
        profile.SCOPE_SHEET,
    ]


def test_the_headers_are_the_profile_in_order():
    book = load_workbook(io.BytesIO(template.workbook()))
    for sheet in Sheet:
        # Trailing empties: the worked example writes a note beside itself, and
        # that widens the sheet without adding a column.
        header = [cell.value for cell in book[str(sheet)][1] if cell.value is not None]
        assert header == [item.name for item in profile.columns_for(sheet)]


def test_the_guidance_says_what_each_blank_does():
    """The single most damaging thing a person can do is guess into a cell."""
    book = load_workbook(io.BytesIO(template.workbook()))
    text = "\n".join(
        str(cell.value)
        for row in book[profile.GUIDE_SHEET].iter_rows()
        for cell in row
        if cell.value
    )
    assert "never 'treat it as zero'" in text
    assert "An assumption fills it" in text
    assert "Derived from the policy total" in text
    assert profile.PROFILE_VERSION in text


def test_the_blank_template_opens_with_a_worked_row():
    """A person reads the example before the explanation.

    Shaded, and marked in a note beside it rather than in a cell a column
    needs, so every column still shows its own example. The reader drops it
    whether or not they delete it.
    """
    book = load_workbook(io.BytesIO(template.workbook()))
    sheet = book[profile.RISK_SHEET]
    row = [str(cell.value) for cell in sheet[2]]

    # Every column shows its own example, and the row says beside itself that
    # it is an example.
    assert "2026_06_PFAC8716" in row
    assert "QEQ" in row
    assert any(value.startswith("EXAMPLE ROW") for value in row)


def test_the_example_row_never_becomes_exposure():
    """Left in place by a hurried user, it must not be read as a risk."""
    book = load_workbook(io.BytesIO(template.workbook()))
    buffer = io.BytesIO()
    book.save(buffer)

    result = intake.read_workbook(io.BytesIO(buffer.getvalue()))

    assert result.risks.rows == []
    assert result.policies.rows == []


def test_every_column_shows_an_example_in_the_guidance():
    book = load_workbook(io.BytesIO(template.workbook()))
    guidance = " | ".join(
        str(cell.value)
        for row in book[profile.GUIDE_SHEET].iter_rows()
        for cell in row
        if cell.value
    )

    for column in profile.COLUMNS:
        assert column.example in guidance, column.name


def test_the_policy_id_is_described_as_the_premium_system_writes_it():
    """"Account number" was ambiguous; the premium system's policy id is not."""
    column = profile.column("Policy ID")

    assert column.oed_field == "AccNumber"
    assert column.example == "2026_06_PFAC8716"
    assert "underwriting year" in column.help_text


def test_the_guidance_names_the_fields_cass_fills_in_itself():
    book = load_workbook(io.BytesIO(template.workbook()))
    text = "\n".join(
        str(cell.value)
        for row in book[profile.GUIDE_SHEET].iter_rows()
        for cell in row
        if cell.value
    )
    assert "PortNumber" in text
    assert "the template does not ask for them" in text


def test_the_same_generator_produces_a_populated_file():
    """A template and an export that disagreed would be two formats, one name."""
    book = load_workbook(io.BytesIO(template.workbook(risks=[risk()])))
    sheet = book[profile.RISK_SHEET]
    row = dict(
        zip(
            [cell.value for cell in sheet[1]],
            [cell.value for cell in sheet[2]],
            strict=True,
        )
    )
    assert row["Policy ID"] == "A1"
    assert row["Building value"] == "1000000.00"


# -- reading one back ---------------------------------------------------------------------

def test_a_completed_template_reads_with_no_findings():
    result = read([risk()], [policy()])
    assert result.is_readable
    assert result.findings == []
    assert len(result.risks.rows) == 1


def test_values_are_typed_the_way_the_model_needs_them():
    result = read([risk(**{"Occupancy": "1100", "Storeys": "4"})], [policy()])
    row = result.risks.rows[0]
    assert row.get("Building value") == Decimal("1000000.00")
    assert row.get("Latitude") == Decimal("-6.2088")
    assert row.get("Storeys") == 4
    assert row.get("Occupancy") == "1100"


def test_money_keeps_its_cents_exactly():
    """A float would lose the reconciliation the whole platform rests on."""
    result = read([risk(**{"Building value": "1234567.89"})], [policy()])
    assert result.risks.rows[0].get("Building value") == Decimal("1234567.89")


def test_thousands_separators_are_tolerated():
    result = read([risk(**{"Building value": "1,234,567.89"})], [policy()])
    assert result.risks.rows[0].get("Building value") == Decimal("1234567.89")


def test_a_share_outside_zero_to_one_says_how_to_write_it():
    result = read([risk()], [policy(**{"Signed share": "15"})])
    finding = next(item for item in result.findings if item.field == "Signed share")
    assert "written as 0.15" in finding.message


def test_a_missing_required_value_is_reported_against_its_row():
    result = read([risk(**{"Country": ""})], [policy()])
    finding = next(item for item in result.findings if item.code == "missing_value")
    assert finding.field == "Country"
    assert finding.row_number == 2


def test_a_missing_required_column_is_fatal():
    payload = template.workbook(risks=[risk()])
    book = load_workbook(io.BytesIO(payload))
    sheet = book[profile.RISK_SHEET]
    for cell in sheet[1]:
        if cell.value == "Country":
            cell.value = None
    buffer = io.BytesIO()
    book.save(buffer)

    result = intake.read_workbook(io.BytesIO(buffer.getvalue()))
    assert result.is_readable is False
    assert "Country" in result.risks.missing_columns


def test_an_unrecognised_column_is_reported_but_not_fatal():
    """A source system adding a field it did not have before is normal."""
    payload = template.workbook(risks=[risk()])
    book = load_workbook(io.BytesIO(payload))
    sheet = book[profile.RISK_SHEET]
    sheet.cell(row=1, column=sheet.max_column + 1, value="Underwriter notes")
    buffer = io.BytesIO()
    book.save(buffer)

    result = intake.read_workbook(io.BytesIO(buffer.getvalue()))
    assert result.is_readable is True
    assert "Underwriter notes" in result.risks.unrecognised_columns


def test_a_finding_quotes_the_cell_so_a_person_can_correct_it():
    """Blanking values would make a findings list unactionable."""
    result = read([risk(**{"Storeys": "four"})])
    finding = next(item for item in result.findings if item.field == "Storeys")
    assert finding.value == "four"


def test_a_workbook_that_is_not_a_template_says_so():
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "Sheet1"
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(intake.IntakeError, match="not a CASS intake template"):
        intake.read_workbook(io.BytesIO(buffer.getvalue()))


def test_a_trailing_empty_row_is_how_a_spreadsheet_ends():
    payload = template.workbook(risks=[risk()])
    book = load_workbook(io.BytesIO(payload))
    book[profile.RISK_SHEET].cell(row=5, column=1, value=None)
    buffer = io.BytesIO()
    book.save(buffer)

    result = intake.read_workbook(io.BytesIO(buffer.getvalue()))
    assert len(result.risks.rows) == 1


# -- blank is an answer ---------------------------------------------------------------------

def test_a_risk_stating_no_value_defers_to_the_allocation():
    """The case the old two-sheet workbook forced on every policy."""
    unvalued = risk(**{"Building value": ""})
    result = read([unvalued], [policy(**{"Total insured value": "1000000.00"})])
    assert result.findings == []
    assert len(result.deferred()) == 1
    assert dict(intake.coverage_evidence(result)) == {
        "risks": 1,
        "coverages_stated": 0,
        "risk_total_stated": 0,
        "allocated_from_policy": 1,
    }


def test_a_risk_stating_some_values_has_stated_them_all():
    """One filled cell means the person knew the schedule."""
    partial = risk(**{"Building value": "", "Contents value": "50000.00"})
    result = read([partial], [policy()])
    assert result.deferred() == []
    assert intake.coverage_evidence(result)["coverages_stated"] == 1


def test_a_risk_total_without_a_breakdown_is_its_own_tier():
    """Knowing what a site is worth and not how it splits is the ordinary case.

    It must not be confused with knowing neither: no location allocation
    applies here, only a component split, and recording an allocation
    assumption that never ran would misdescribe the version.
    """
    known = risk(**{"Building value": "", "Total insured value": "1500000.00"})
    result = read([known], [policy()])
    assert result.findings == []
    assert result.deferred() == []
    assert dict(intake.coverage_evidence(result)) == {
        "risks": 1,
        "coverages_stated": 0,
        "risk_total_stated": 1,
        "allocated_from_policy": 0,
    }


def test_stated_coverages_outrank_a_risk_total():
    """Both filled is not a contradiction; the finer evidence simply wins."""
    both = risk(**{"Building value": "900000.00", "Total insured value": "1000000.00"})
    result = read([both], [policy()])
    assert intake.states_coverages(result.risks.rows[0]) is True
    assert intake.coverage_evidence(result)["coverages_stated"] == 1


def test_a_risk_total_needs_no_policy_total_to_fall_back_on():
    """The middle tier is self-sufficient, so it raises no allocation finding."""
    known = risk(**{"Building value": "", "Total insured value": "1500000.00"})
    result = read([known])
    assert [item for item in result.findings if item.code == "no_value_to_allocate"] == []


def test_a_risk_with_nothing_to_allocate_is_reported_not_valued_at_nothing():
    result = read([risk(**{"Building value": ""})], [policy()])
    finding = next(item for item in result.findings if item.code == "no_value_to_allocate")
    assert "A risk worth nothing is not what an empty row means" in finding.message


def test_a_portfolio_with_no_policy_sheet_still_reads():
    """Policy terms are optional by design; ground-up loss does not need them."""
    result = read([risk()])
    assert result.is_readable is True
    assert result.has_policy_terms is False
    assert result.findings == []


# -- the join is checked, not inferred ----------------------------------------------------------

def test_a_policy_covering_no_risk_is_reported():
    result = read([risk()], [policy(), policy(**{"Policy ID": "A9"})])
    finding = next(item for item in result.findings if item.code == "policy_without_risks")
    assert "A9" in finding.message


def test_a_risk_with_no_policy_is_reported_where_others_have_one():
    result = read(
        [risk(), risk(**{"Policy ID": "A2", "Risk reference": "1"})],
        [policy()],
    )
    finding = next(item for item in result.findings if item.code == "risk_without_policy")
    assert "A2" in finding.message
    assert "Ground-up loss is unaffected" in finding.message


def test_the_same_risk_reference_twice_in_one_account_is_refused():
    result = read([risk(), risk()], [policy()])
    finding = next(item for item in result.findings if item.code == "duplicate_risk")
    assert "put the same building in twice" in finding.message


def test_the_same_risk_reference_in_two_accounts_is_fine():
    """A location number is unique within an account, not across a portfolio."""
    result = read(
        [risk(), risk(**{"Policy ID": "A2"})],
        [policy(), policy(**{"Policy ID": "A2"})],
    )
    assert [item for item in result.findings if item.code == "duplicate_risk"] == []


def test_a_repeated_policy_layer_is_refused():
    result = read([risk()], [policy(**{"Layer": "1"}), policy(**{"Layer": "1"})])
    assert any(item.code == "duplicate_policy" for item in result.findings)


def test_two_layers_of_one_policy_are_fine():
    result = read([risk()], [policy(**{"Layer": "1"}), policy(**{"Layer": "2"})])
    assert [item for item in result.findings if item.code == "duplicate_policy"] == []


# -- deductibles ---------------------------------------------------------------------------------

def test_a_single_risk_deductible_on_both_sheets_is_reported_with_its_total():
    """The same 25,000 typed on both sheets is charged twice, and the finding says so."""
    result = read(
        [risk(**{"Risk deductible": "25000"})],
        [policy(**{"Policy deductible": "25000"})],
    )
    finding = next(
        item for item in result.findings if item.code == "deductible_on_both_sheets"
    )
    assert finding.sheet == profile.RISK_SHEET
    assert "1 single-risk policy states" in finding.message
    assert "25,000 + 25,000 = 50,000" in finding.message
    assert "keep it on the Policies sheet" in finding.message


def test_single_risk_policies_with_both_deductibles_are_reported_once_with_a_count():
    """Most of a book is one risk per policy, so one finding, not one per policy."""
    result = read(
        [
            risk(**{"Risk deductible": "25000"}),
            risk(**{"Policy ID": "A2", "Risk deductible": "10000"}),
        ],
        [
            policy(**{"Policy deductible": "25000"}),
            policy(**{"Policy ID": "A2", "Policy deductible": "10000"}),
        ],
    )
    findings = [item for item in result.findings if item.code == "deductible_on_both_sheets"]
    assert len(findings) == 1
    assert findings[0].value == "2"
    assert "A1, A2" in findings[0].message


def test_a_deductible_on_one_sheet_only_is_not_reported():
    result = read([risk()], [policy(**{"Policy deductible": "25000"})])
    assert [item for item in result.findings if item.code == "deductible_on_both_sheets"] == []


def test_a_multi_risk_policy_may_carry_both_deductibles():
    """Per-site deductibles and a policy deductible are two different terms there."""
    result = read(
        [
            risk(**{"Risk deductible": "25000"}),
            risk(**{"Risk reference": "2", "Risk deductible": "25000"}),
        ],
        [policy(**{"Policy deductible": "100000"})],
    )
    assert [item for item in result.findings if item.code == "deductible_on_both_sheets"] == []


def test_the_guidance_says_where_a_single_risk_deductible_goes():
    book = load_workbook(io.BytesIO(template.workbook()))
    text = " ".join(
        str(cell.value) for row in book[profile.GUIDE_SHEET].iter_rows() for cell in row
        if cell.value
    )
    assert "put the deductible and limit on the Policies sheet only" in text
    assert "Repeat the policy deductible and limit on every layer's row" in text


# -- lineage ------------------------------------------------------------------------------------

def test_a_read_records_the_profile_and_parser_that_produced_it():
    result = read([risk()], [policy()])
    assert result.profile_version == profile.PROFILE_VERSION
    assert result.parser_version == intake.PARSER_VERSION


def test_a_date_is_read_from_either_a_date_cell_or_iso_text():
    result = read([risk()], [policy(**{"Inception date": "2026-01-01"})])
    assert result.policies.rows[0].get("Inception date") == dt.date(2026, 1, 1)


@pytest.mark.parametrize("value", ["Yes", "yes", "TRUE", "1"])
def test_a_flag_accepts_the_ways_people_write_it(value):
    result = read([risk(**{"Needs review": value})], [policy()])
    assert result.risks.rows[0].get("Needs review") is True


def test_an_unreadable_flag_says_what_is_expected():
    result = read([risk(**{"Needs review": "maybe"})], [policy()])
    finding = next(item for item in result.findings if item.field == "Needs review")
    assert "Yes or No" in finding.message


def test_the_profile_is_serialisable_for_the_schema_endpoint():
    payload = profile.as_dict()
    assert payload["profile_version"] == profile.PROFILE_VERSION
    assert set(payload["sheets"]) == {
        profile.RISK_SHEET,
        profile.POLICY_SHEET,
        profile.CONTRACT_SHEET,
        profile.SCOPE_SHEET,
    }
    first = payload["sheets"][profile.RISK_SHEET][0]
    assert first["oed_field"] == "AccNumber"
    assert first["reads_as"] == str(DataType.TEXT)
    assert payload["oed_fields_not_requested"]["location"] == ["IsTenant", "LocGroup"]


# -- policy terms and reinsurance ----------------------------------------------------------------

def codes(result) -> list[str]:
    return [item.code for item in result.findings]


def test_the_reinsurance_sheets_open_with_a_layered_programme():
    """One example row cannot show two layers sharing a contract number."""
    book = load_workbook(io.BytesIO(template.workbook()))
    contracts = [
        [cell.value for cell in row]
        for row in book[profile.CONTRACT_SHEET].iter_rows(min_row=2)
        if any(cell.value for cell in row)
    ]
    assert len(contracts) == 3
    assert [row[0] for row in contracts] == ["1", "2", "2"]
    assert all(row[-1] == template.EXAMPLE_MARKER for row in contracts)


def test_the_worked_reinsurance_rows_never_become_contracts():
    result = intake.read_workbook(io.BytesIO(template.workbook()))
    assert result.contracts.rows == []
    assert result.scope.rows == []


def test_a_workbook_from_before_the_reinsurance_sheets_still_reads():
    book = load_workbook(io.BytesIO(template.workbook(risks=[risk()], policies=[policy()])))
    for name in (profile.CONTRACT_SHEET, profile.SCOPE_SHEET):
        del book[name]
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    result = intake.read_workbook(buffer)
    assert result.is_readable
    assert result.contracts.rows == []


def test_contracts_and_scope_become_canonical_records():
    result = read(
        [risk()],
        [policy()],
        [contract(), contract(**{"Layer": "2", "Attachment per event": "25000000"})],
        [{"Contract number": "2"}],
    )
    contracts, scope = intake.reinsurance_records(result)
    assert [(item["contract_number"], item["layer_number"]) for item in contracts] == [(2, 1), (2, 2)]
    assert contracts[1]["occurrence_attachment"] == Decimal("25000000")
    assert scope[0]["business_id"] is None


def test_a_contract_layer_given_twice_is_reported():
    result = read([risk()], [policy()], [contract(), contract()], [{"Contract number": "2"}])
    assert "duplicate_contract_layer" in codes(result)


def test_layers_of_one_contract_must_be_the_same_kind_of_contract():
    result = read(
        [risk()],
        [policy()],
        [contract(), contract(**{"Layer": "2", "Inuring priority": "3"})],
        [{"Contract number": "2"}],
    )
    assert "contract_layers_disagree" in codes(result)


def test_scope_and_contracts_must_name_each_other():
    result = read([risk()], [policy()], [contract()], [{"Contract number": "7"}])
    assert "scope_without_contract" in codes(result)
    assert "contract_without_scope" in codes(result)


def test_a_policy_filled_in_for_its_value_alone_is_not_a_problem():
    """A ground-up book fills the Policies sheet for the join and the allocation."""
    result = read([risk()], [policy(**{"Total insured value": "1000000"})])
    assert "policy_terms_incomplete" not in codes(result)


def test_terms_without_a_layer_attachment_and_limit_are_reported():
    result = read([risk()], [policy(**{"Policy deductible": "25000"})])
    finding = next(item for item in result.findings if item.code == "policy_terms_incomplete")
    assert "Write 0 as the attachment" in finding.message


def test_complete_accounts_need_every_row_complete():
    records = [
        {"business_id": "A1", "layer_attachment": "0", "layer_limit": "100"},
        {"business_id": "A2", "layer_attachment": "0", "layer_limit": "100"},
        {"business_id": "A2", "layer_attachment": "100", "layer_limit": None},
    ]
    assert intake.complete_accounts(records) == {"A1"}


def test_layers_stating_different_deductibles_are_reported():
    result = read(
        [risk()],
        [
            policy(**{"Layer": "1", "Layer attachment": "0", "Layer limit": "3000000",
                      "Policy deductible": "100000"}),
            policy(**{"Layer": "2", "Layer attachment": "3000000", "Layer limit": "7000000"}),
        ],
    )
    finding = next(item for item in result.findings if item.code == "layer_terms_differ")
    assert "A1 P1" in finding.message
