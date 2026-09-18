"""Reading a completed CASS intake template.

The counterpart of ``template``: same profile, read instead of written. What
changes most against the previous intake shape is not the parsing but the join.
The account reference is on both sheets because a person put it there, so this
module *checks* the join rather than reconstructing it -- an orphan is a
mismatch between two stated facts, not a guess that failed.

What a blank means is the whole point of the format, and value has two
separate unknowns that the reader is careful not to conflate: *which risk holds
the money*, and *how that money splits between building, contents and the
rest*. A person can easily know the first and not the second, and a format that
forced them to answer both would collect a guess for one of them.

So value reads in three tiers, best evidence first.

**Coverage columns filled.** The person knew the schedule. Empty coverage cells
on such a row are zero at that coverage, not unknown -- reading them as unknown
would let a single contents figure drag a whole building value into an
assumption nobody asked for.

**Coverage blank, risk total filled.** The value at this risk is known; its
breakdown is not. A named component assumption divides it, and the version
records which one. No location allocation is involved and none is recorded.

**Both blank.** Neither is known, which is a legitimate answer and the one the
previous workbook forced on every policy. The policy's total is divided across
its risks under a selectable allocation scenario, then split by coverage. The
reader reports a risk in this tier whose policy states no total either, rather
than letting it arrive worth nothing.

**A required cell is never assumed.** Identity, geography, currency and perils
are refused rather than filled, because there is no honest default for where a
building is.
"""

from __future__ import annotations

import dataclasses
import datetime as dt
from collections.abc import Iterator, Mapping, Sequence
from decimal import Decimal, InvalidOperation
from typing import Any, BinaryIO

from cass_oed.schema import DataType

from . import profile
from .profile import PROFILE_VERSION, Column, Sheet
from .records import ExtractReadError, Finding, SheetRead, SourceRow

#: Bumped when this reader's interpretation changes, independently of the
#: profile: the same file read by a later parser should be traceable to which.
PARSER_VERSION = "2.1.0"

#: The coverage columns a risk row may carry.
COVERAGE_COLUMNS = (
    "Building value",
    "Other structures value",
    "Contents value",
    "Business interruption value",
)

#: The risk-level total, for a schedule that knows what a site is worth but not
#: how the worth divides.
RISK_TOTAL_COLUMN = "Total insured value"


class IntakeError(ExtractReadError):
    """Raised when the workbook cannot be read as a CASS intake template."""


@dataclasses.dataclass(slots=True)
class IntakeRead:
    """Every sheet of a completed template, parsed and cross-checked."""

    risks: SheetRead
    policies: SheetRead
    #: Problems that span the sheets: a stated join that does not hold, a risk
    #: deferring to an allocation its policy cannot supply.
    cross_findings: list[Finding] = dataclasses.field(default_factory=list)
    #: The reinsurance sheets. Optional, and absent from a 1.0.0 workbook.
    contracts: SheetRead = dataclasses.field(default_factory=lambda: _empty(Sheet.CONTRACT))
    scope: SheetRead = dataclasses.field(default_factory=lambda: _empty(Sheet.SCOPE))
    profile_version: str = PROFILE_VERSION
    parser_version: str = PARSER_VERSION

    @property
    def sheets(self) -> tuple[SheetRead, ...]:
        return (self.risks, self.policies, self.contracts, self.scope)

    @property
    def findings(self) -> list[Finding]:
        return [*(item for sheet in self.sheets for item in sheet.findings), *self.cross_findings]

    @property
    def is_readable(self) -> bool:
        """Whether the shape is close enough to the contract to go on.

        A missing required column is fatal; a Policies or reinsurance sheet
        being absent entirely is not, because both are optional by design.
        """
        return not any(sheet.missing_columns for sheet in self.sheets)

    @property
    def has_policy_terms(self) -> bool:
        return bool(self.policies.rows)

    def as_dict(self) -> dict[str, Any]:
        return {
            "profile_version": self.profile_version,
            "parser_version": self.parser_version,
            "risk_count": len(self.risks.rows),
            "policy_count": len(self.policies.rows),
            "contract_count": len(self.contracts.rows),
            "scope_count": len(self.scope.rows),
            "accounts": len(self.accounts()),
            "readable": self.is_readable,
            "has_policy_terms": self.has_policy_terms,
            "finding_count": len(self.findings),
            "risks_deferring_to_allocation": len(self.deferred()),
        }

    def accounts(self) -> set[str]:
        return {
            str(row.get("Policy ID") or "").strip()
            for row in self.risks.rows
            if row.get("Policy ID")
        }

    def deferred(self) -> list[SourceRow]:
        """Risk rows that state no value at all and so await an allocation."""
        return [row for row in self.risks.rows if defers_to_allocation(row)]


def states_coverages(row: SourceRow) -> bool:
    """Whether the row states at least one coverage value outright."""
    return any(row.values.get(column) is not None for column in COVERAGE_COLUMNS)


def states_risk_total(row: SourceRow) -> bool:
    """Whether the row states what the risk is worth without breaking it down."""
    return row.values.get(RISK_TOTAL_COLUMN) is not None


def defers_to_allocation(row: SourceRow) -> bool:
    """Whether the row states no value at all, at either level."""
    return not states_coverages(row) and not states_risk_total(row)


def read_workbook(source: BinaryIO | str) -> IntakeRead:
    """Read a completed CASS intake workbook."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - packaging error
        raise IntakeError("openpyxl is required to read a template.") from exc

    try:
        book = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        # Anything openpyxl cannot open: a CSV renamed, a truncated upload, a
        # PDF. The source is still registered; this is what the batch records.
        raise IntakeError(
            "The file could not be opened as a workbook, so it is not a CASS "
            f"intake template. ({exc})"
        ) from exc
    names = {name.strip().lower(): name for name in book.sheetnames}
    if profile.RISK_SHEET.lower() not in names:
        raise IntakeError(
            f"The workbook has no {profile.RISK_SHEET!r} sheet, so it is not a CASS "
            "intake template. Sheets present: " + ", ".join(book.sheetnames) + "."
        )

    risks = _read_sheet(book[names[profile.RISK_SHEET.lower()]], Sheet.RISK)
    def optional(sheet: Sheet) -> SheetRead:
        name = names.get(str(sheet).lower())
        return _read_sheet(book[name], sheet) if name else _empty(sheet)

    policies = optional(Sheet.POLICY)
    contracts = optional(Sheet.CONTRACT)
    scope = optional(Sheet.SCOPE)
    book.close()

    read = IntakeRead(risks=risks, policies=policies, contracts=contracts, scope=scope)
    read.cross_findings = list(cross_check(read))
    return read


def read_rows(
    sheet: Sheet, columns: Sequence[str], rows: Sequence[Sequence[Any]]
) -> SheetRead:
    """Read one sheet from values, for callers that already have them."""
    return _read_values(sheet, list(columns), rows)


def _empty(sheet: Sheet) -> SheetRead:
    return SheetRead(
        sheet=str(sheet),
        columns=(),
        rows=[],
        missing_columns=(),
        unrecognised_columns=(),
        findings=[],
    )


def _is_example(cells: Sequence[Any]) -> bool:
    """Whether this row is the template's own worked example.

    Read from the whole row rather than the mapped columns: the marker sits
    beside the example, in a cell no column claims, so that every column of the
    example can show its own value.
    """
    return any(
        cell is not None and str(cell).strip().upper().startswith("EXAMPLE ROW")
        for cell in cells
    )


def _read_sheet(worksheet, sheet: Sheet) -> SheetRead:
    iterator = worksheet.iter_rows(values_only=True)
    try:
        header = next(iterator)
    except StopIteration:
        return _empty(sheet)
    columns = [str(cell).strip() if cell is not None else "" for cell in header]
    return _read_values(sheet, columns, iterator)


def _read_values(
    sheet: Sheet, columns: list[str], rows: Any
) -> SheetRead:
    expected = profile.columns_for(sheet)
    by_name = {item.name.strip().lower(): item for item in expected}
    resolved = [by_name.get(name.strip().lower()) for name in columns]

    present = {item.name for item in resolved if item is not None}
    missing = tuple(
        item.name for item in expected if item.required and item.name not in present
    )
    unrecognised = tuple(
        name
        for name, item in zip(columns, resolved, strict=True)
        if name and item is None
    )

    findings: list[Finding] = []
    parsed: list[SourceRow] = []

    # Row 1 is the header, so the first data row is row 2.
    for offset, cells in enumerate(rows, start=2):
        raw: dict[str, str] = {}
        values: dict[str, Any] = {}
        blank = True

        for position, item in enumerate(resolved):
            if item is None:
                continue
            cell = cells[position] if position < len(cells) else None
            text = "" if cell is None else str(cell).strip()
            raw[item.name] = text
            if text:
                blank = False
            value, finding = _coerce(sheet, offset, item, cell)
            if finding is not None:
                findings.append(finding)
            values[item.name] = value

        if blank:
            # A trailing empty row is how a spreadsheet ends, not a problem.
            continue

        if _is_example(cells):
            # The template ships a worked row so a person can see what a filled
            # in one looks like. It says in its first cell that it should go,
            # and it is dropped here whether or not they deleted it: a hurried
            # user must not be able to turn the example into exposure.
            continue

        for item in expected:
            if item.required and values.get(item.name) in (None, ""):
                findings.append(
                    Finding(
                        sheet=str(sheet),
                        row_number=offset,
                        field=item.name,
                        code="missing_value",
                        message=f"{item.name} is required and was not supplied.",
                    )
                )

        parsed.append(SourceRow(sheet=str(sheet), row_number=offset, raw=raw, values=values))

    return SheetRead(
        sheet=str(sheet),
        columns=tuple(columns),
        rows=parsed,
        missing_columns=missing,
        unrecognised_columns=unrecognised,
        findings=findings,
    )


def _coerce(
    sheet: Sheet, row_number: int, item: Column, cell: Any
) -> tuple[Any, Finding | None]:
    """Turn one cell into its typed value, or explain why it could not be."""
    if cell is None or (isinstance(cell, str) and not cell.strip()):
        return None, None

    def bad(message: str) -> tuple[None, Finding]:
        return None, Finding(
            sheet=str(sheet),
            row_number=row_number,
            field=item.name,
            code="unreadable_value",
            message=message,
            # The cell's own text, so a person can see what needs correcting.
            # Nothing in a Klapton Re portfolio is withheld from a Klapton Re
            # colleague; what a finding must not become is a route for whole
            # source rows into a log, and quoting one cell is not that.
            value=str(cell)[:120],
        )

    text = str(cell).strip()

    match item.reads_as:
        case DataType.TEXT | DataType.PERIL | DataType.CURRENCY:
            return text, None

        case DataType.INTEGER:
            try:
                return int(Decimal(text)), None
            except (InvalidOperation, TypeError, ValueError):
                return bad(f"{item.name} is not a whole number.")

        case DataType.MONEY | DataType.DECIMAL:
            # Decimal from the string form: exact reconciliation cannot survive
            # a float, and a spreadsheet hands out floats freely.
            try:
                value = Decimal(text.replace(",", ""))
            except (InvalidOperation, TypeError, ValueError):
                return bad(f"{item.name} is not an amount.")
            if value < 0:
                return bad(f"{item.name} is negative, which cannot be interpreted.")
            return value, None

        case DataType.RATE:
            try:
                value = Decimal(text)
            except (InvalidOperation, TypeError, ValueError):
                return bad(f"{item.name} is not a proportion.")
            if not 0 <= value <= 1:
                return bad(
                    f"{item.name} of {value} is not between 0 and 1. A share of 15% "
                    "is written as 0.15."
                )
            return value, None

        case DataType.LATITUDE | DataType.LONGITUDE:
            try:
                value = Decimal(text)
            except (InvalidOperation, TypeError, ValueError):
                return bad(f"{item.name} is not a coordinate.")
            limit = Decimal(90) if item.reads_as is DataType.LATITUDE else Decimal(180)
            if not -limit <= value <= limit:
                return bad(f"{item.name} of {value} is outside the valid range.")
            return value, None

        case DataType.DATE:
            if isinstance(cell, dt.datetime):
                return cell.date(), None
            if isinstance(cell, dt.date):
                return cell, None
            try:
                return dt.date.fromisoformat(text), None
            except ValueError:
                return bad(f"{item.name} is not a date. Use YYYY-MM-DD.")

        case DataType.FLAG:
            lowered = text.lower()
            if lowered in ("yes", "y", "true", "1"):
                return True, None
            if lowered in ("no", "n", "false", "0"):
                return False, None
            return bad(f"{item.name} should be Yes or No.")

    return text, None  # pragma: no cover - every DataType is handled above


def cross_check(read: IntakeRead) -> Iterator[Finding]:
    """Problems that only appear when the sheets are read together.

    This is what replaces the old join report, and it is much shorter for one
    reason: the account reference was supplied rather than inferred, so a
    mismatch is a contradiction between two stated facts instead of a
    reconstruction that did not converge.
    """
    yield from _duplicate_risks(read.risks)
    yield from _duplicate_policies(read.policies)
    yield from _orphans(read)
    yield from _allocation_needs(read)
    yield from _construction_without_occupancy(read.risks)
    yield from _deductible_on_both_sheets(read)
    yield from _incomplete_policy_terms(read)
    yield from _layers_that_disagree(read)
    yield from _contract_structure(read)


def _duplicate_risks(sheet: SheetRead) -> Iterator[Finding]:
    seen: dict[tuple[str, str], int] = {}
    for row in sheet.rows:
        key = (
            str(row.get("Policy ID") or "").strip(),
            str(row.get("Risk reference") or "").strip(),
        )
        if not all(key):
            continue
        if key in seen:
            yield Finding(
                sheet=sheet.sheet,
                row_number=row.row_number,
                field="Risk reference",
                code="duplicate_risk",
                message=(
                    f"Risk {key[1]} of account {key[0]} is already on row {seen[key]}. "
                    "A risk reference identifies one property within its account, so "
                    "two rows sharing one would put the same building in twice."
                ),
            )
        else:
            seen[key] = row.row_number


def _duplicate_policies(sheet: SheetRead) -> Iterator[Finding]:
    seen: dict[tuple[str, str, str], int] = {}
    for row in sheet.rows:
        layer = row.get("Layer")
        key = (
            str(row.get("Policy ID") or "").strip(),
            str(row.get("Policy reference") or "").strip(),
            "" if layer is None else str(layer),
        )
        if not (key[0] and key[1]):
            continue
        if key in seen:
            yield Finding(
                sheet=sheet.sheet,
                row_number=row.row_number,
                field="Policy reference",
                code="duplicate_policy",
                message=(
                    f"Policy {key[1]} of account {key[0]}"
                    + (f" layer {key[2]}" if key[2] else "")
                    + f" is already on row {seen[key]}."
                ),
            )
        else:
            seen[key] = row.row_number


def _accounts(count: int) -> str:
    return "account has" if count == 1 else "accounts have"


def _orphans(read: IntakeRead) -> Iterator[Finding]:
    """Accounts named on one sheet and not the other.

    Aggregated deliberately. A book where most accounts are not yet geocoded
    is an ordinary state of affairs, and reporting it once with a count is
    something a person can act on; reporting it a thousand times is noise that
    buries the findings that mean something.
    """
    risk_accounts = read.accounts()
    policy_accounts: dict[str, int] = {}
    for row in read.policies.rows:
        account = str(row.get("Policy ID") or "").strip()
        if account:
            policy_accounts.setdefault(account, row.row_number)

    uncovered = sorted(
        (account, row_number)
        for account, row_number in policy_accounts.items()
        if account not in risk_accounts
    )
    if uncovered:
        yield Finding(
            sheet=read.policies.sheet,
            row_number=uncovered[0][1],
            field="Policy ID",
            code="policy_without_risks",
            message=(
                f"{len(uncovered)} {_accounts(len(uncovered))} policy terms but "
                "no risks, so nothing would be modelled under them. This is "
                "expected where a book is only partly geocoded. First: "
                + ", ".join(account for account, _ in uncovered[:5])
                + "."
            ),
            value=str(len(uncovered)),
        )

    if not policy_accounts:
        return
    unpriced = sorted(
        {
            str(row.get("Policy ID") or "").strip(): row.row_number
            for row in read.risks.rows
            if str(row.get("Policy ID") or "").strip()
            and str(row.get("Policy ID") or "").strip() not in policy_accounts
        }.items()
    )
    if unpriced:
        yield Finding(
            sheet=read.risks.sheet,
            row_number=unpriced[0][1],
            field="Policy ID",
            code="risk_without_policy",
            message=(
                f"{len(unpriced)} {_accounts(len(unpriced))} risks but no policy "
                "terms, while others have them. Ground-up loss is unaffected; "
                "insured loss cannot be calculated for them. First: "
                + ", ".join(account for account, _ in unpriced[:5])
                + "."
            ),
            value=str(len(unpriced)),
        )


def _construction_without_occupancy(sheet: SheetRead) -> Iterator[Finding]:
    """A construction code that cannot reach a vulnerability function.

    OED routes on the pair, and the occupancy is the half that decides which
    table is consulted, so a construction on its own reaches nothing. The
    occupancy assumption would fill both -- silently discarding the one real
    attribute the schedule stated. Better to say so while someone can still
    supply the occupancy it belongs with.
    """
    for row in sheet.rows:
        construction = str(row.get("Construction") or "").strip()
        occupancy = str(row.get("Occupancy") or "").strip()
        if not construction or occupancy:
            continue
        yield Finding(
            sheet=sheet.sheet,
            row_number=row.row_number,
            field="Occupancy",
            code="construction_without_occupancy",
            message=(
                f"This risk states a construction code with no occupancy code. OED "
                f"requires an occupancy, and a construction alone cannot reach a "
                f"vulnerability function -- so {construction} would be discarded and "
                "the occupancy assumption would supply both. Supply the occupancy."
            ),
        )


def _deductible_on_both_sheets(read: IntakeRead) -> Iterator[Finding]:
    """Single-risk policies whose deductible is written on both sheets.

    A risk deductible comes off first and the policy deductible comes off what
    is left. Over several risks those are two different terms; over one, it is
    usually one deductible typed twice, and it would be charged twice.

    Aggregated like the orphans: most of a facultative book is one risk per
    policy, and a source system that fills both columns fills them on every row.
    """
    risks: dict[str, list[SourceRow]] = {}
    for row in read.risks.rows:
        account = str(row.get("Policy ID") or "").strip()
        if account:
            risks.setdefault(account, []).append(row)
    policy_deductible: dict[str, Decimal] = {}
    for row in read.policies.rows:
        account = str(row.get("Policy ID") or "").strip()
        amount = row.get("Policy deductible")
        if account and amount is not None and amount > 0:
            policy_deductible.setdefault(account, amount)

    doubled = sorted(
        (account, rows[0].row_number, rows[0].get("Risk deductible"), policy_deductible[account])
        for account, rows in risks.items()
        if len(rows) == 1
        and account in policy_deductible
        and rows[0].get("Risk deductible") is not None
        and rows[0].get("Risk deductible") > 0
    )
    if not doubled:
        return
    account, row_number, at_risk, at_policy = doubled[0]
    policies = "policy states" if len(doubled) == 1 else "policies state"
    yield Finding(
        sheet=read.risks.sheet,
        row_number=row_number,
        field="Risk deductible",
        code="deductible_on_both_sheets",
        message=(
            f"{len(doubled)} single-risk {policies} a deductible on both sheets. "
            "The risk deductible comes off first and the policy deductible comes off "
            f"what is left, so both would be charged: for {account}, {at_risk:,} + "
            f"{at_policy:,} = {at_risk + at_policy:,}. If they are the same deductible, "
            "keep it on the Policies sheet and clear Risk deductible. First: "
            + ", ".join(item[0] for item in doubled[:5])
            + "."
        ),
        value=str(len(doubled)),
    )


def _text(value: Any) -> str:
    return str(value or "").strip()


def terms_complete(record: Mapping[str, Any]) -> bool:
    """Whether one policy row states enough to calculate insured loss.

    A layer attachment and a layer limit, as the profile says. Read from a
    canonical record, so promotion and the reader apply one rule.
    """
    return all(
        record.get(name) not in (None, "") for name in ("layer_attachment", "layer_limit")
    )


def complete_accounts(policies: Sequence[Mapping[str, Any]]) -> set[str]:
    """Policy IDs whose every policy row can be written as OED terms.

    Canonical records in, so the import report and the promotion that writes
    the account file cannot disagree about which accounts have terms.
    """
    rows: dict[str, list[Mapping[str, Any]]] = {}
    for record in policies:
        account = _text(record.get("business_id"))
        if account:
            rows.setdefault(account, []).append(record)
    return {
        account for account, records in rows.items() if all(map(terms_complete, records))
    }


#: Policy columns that say a row is meant to carry terms. A row with none of
#: them is there for the join and the allocation, which is a ground-up book and
#: not a mistake.
TERM_COLUMNS = (
    "Layer",
    "Signed share",
    "Layer limit",
    "Layer attachment",
    "Policy deductible",
    "Policy limit",
)


def _incomplete_policy_terms(read: IntakeRead) -> Iterator[Finding]:
    """Policy IDs that state some terms but not a layer attachment and limit.

    Not an error: promotion writes them as stated, and OED reads the blanks as
    no limit and no attachment. But the insured loss of such a policy is its
    ground-up loss, and the person filling in the workbook should hear that
    now, while they can still supply the terms. Aggregated: a source system
    that leaves a column out leaves it out of every row.
    """
    incomplete: dict[str, int] = {}
    for row in read.policies.rows:
        account = _text(row.get("Policy ID"))
        states_terms = any(row.get(name) is not None for name in TERM_COLUMNS)
        if account and states_terms and not terms_complete(policy_record(row)):
            incomplete.setdefault(account, row.row_number)
    if not incomplete:
        return
    accounts = sorted(incomplete)
    subject = "Policy ID has" if len(accounts) == 1 else "Policy IDs have"
    yield Finding(
        sheet=read.policies.sheet,
        row_number=incomplete[accounts[0]],
        field="Layer attachment",
        code="policy_terms_incomplete",
        message=(
            f"{len(accounts)} {subject} a policy row with no layer attachment or no "
            "layer limit. They are written as stated, so a missing limit leaves the "
            "insured loss equal to the ground-up loss, and CASS flags them as possibly "
            "overstated. Write 0 as the attachment of a layer that pays from the first "
            "loss, and the sum insured as the limit of a share of a whole risk. First: "
            + ", ".join(accounts[:5])
            + "."
        ),
        value=str(len(accounts)),
    )


def _layers_that_disagree(read: IntakeRead) -> Iterator[Finding]:
    """Layers of one policy that state different policy deductibles or limits.

    Each layer reads the policy terms from its own row, so a deductible written
    on the first layer only is not taken off before the second.
    """
    stated: dict[tuple[str, str], dict[str, set[Any]]] = {}
    first_row: dict[tuple[str, str], int] = {}
    for row in read.policies.rows:
        key = (_text(row.get("Policy ID")), _text(row.get("Policy reference")))
        if not all(key):
            continue
        first_row.setdefault(key, row.row_number)
        terms = stated.setdefault(key, {"Policy deductible": set(), "Policy limit": set()})
        for name, values in terms.items():
            values.add(row.get(name))
    differing = sorted(
        key
        for key, terms in stated.items()
        if any(len(values) > 1 for values in terms.values())
    )
    if not differing:
        return
    subject = "policy states" if len(differing) == 1 else "policies state"
    yield Finding(
        sheet=read.policies.sheet,
        row_number=first_row[differing[0]],
        field="Policy deductible",
        code="layer_terms_differ",
        message=(
            f"{len(differing)} layered {subject} a different policy deductible or "
            "limit on different layers. Each layer reads them from its own row, so a "
            "deductible on the first layer only is not taken off before the second. "
            "Repeat them on every layer's row. First: "
            + ", ".join(f"{account} {policy}" for account, policy in differing[:5])
            + "."
        ),
        value=str(len(differing)),
    )


def _contract_structure(read: IntakeRead) -> Iterator[Finding]:
    """What can be said about the reinsurance sheets without the contract rules.

    The rules for each contract type live with the Financial structure screen
    and are applied when the workbook is imported; this is the shape they need
    to be applied to at all -- each layer once, layers of one contract agreeing
    on what kind of contract it is, and scope naming contracts that exist.
    """
    layers: dict[tuple[int, int], int] = {}
    kinds: dict[int, tuple[tuple[str, Any, str], int]] = {}
    for row in read.contracts.rows:
        number = row.get("Contract number")
        if number is None:
            continue
        layer = row.get("Layer") or 1
        if (number, layer) in layers:
            yield Finding(
                sheet=read.contracts.sheet,
                row_number=row.row_number,
                field="Layer",
                code="duplicate_contract_layer",
                message=(
                    f"Contract {number} layer {layer} is already on row "
                    f"{layers[(number, layer)]}. Each layer of a contract is one row."
                ),
            )
            continue
        layers[(number, layer)] = row.row_number
        shape = (
            _text(row.get("Contract type")).upper(),
            row.get("Inuring priority"),
            _text(row.get("Risk level")).upper(),
        )
        if number not in kinds:
            kinds[number] = (shape, row.row_number)
        elif kinds[number][0] != shape:
            yield Finding(
                sheet=read.contracts.sheet,
                row_number=row.row_number,
                field="Contract type",
                code="contract_layers_disagree",
                message=(
                    f"The layers of contract {number} differ in type, inuring "
                    f"priority or risk level from row {kinds[number][1]}. Layers of "
                    "one contract share one scope, so they have to be the same kind "
                    "of contract; give a different contract a different number."
                ),
            )

    scoped: set[int] = set()
    for row in read.scope.rows:
        number = row.get("Contract number")
        if number is None:
            continue
        scoped.add(number)
        if number not in kinds:
            yield Finding(
                sheet=read.scope.sheet,
                row_number=row.row_number,
                field="Contract number",
                code="scope_without_contract",
                message=(
                    f"This row names contract {number}, which the Reinsurance "
                    "contracts sheet does not have."
                ),
            )
    for number, (_, row_number) in sorted(kinds.items()):
        if number not in scoped:
            yield Finding(
                sheet=read.contracts.sheet,
                row_number=row_number,
                field="Contract number",
                code="contract_without_scope",
                message=(
                    f"No Reinsurance scope row names contract {number}, so it covers "
                    "nothing. Add a row with only the contract number to cover the "
                    "whole portfolio."
                ),
            )


def _allocation_needs(read: IntakeRead) -> Iterator[Finding]:
    """Risks that defer to an allocation their policy cannot supply."""
    totals: dict[str, Decimal] = {}
    for row in read.policies.rows:
        account = str(row.get("Policy ID") or "").strip()
        total = row.get("Total insured value")
        if account and total is not None:
            totals[account] = totals.get(account, Decimal("0.00")) + total

    for row in read.risks.rows:
        if not defers_to_allocation(row):
            continue
        account = str(row.get("Policy ID") or "").strip()
        if account and account in totals:
            continue
        yield Finding(
            sheet=read.risks.sheet,
            row_number=row.row_number,
            field=RISK_TOTAL_COLUMN,
            code="no_value_to_allocate",
            message=(
                f"This risk states no value, so it needs its account's total to "
                f"divide -- and account {account or '(unnamed)'} states none either. "
                "Supply the values on the risk, or the total insured value on the "
                "policy. A risk worth nothing is not what an empty row means."
            ),
        )


def coverage_evidence(read: IntakeRead) -> Mapping[str, int]:
    """How many risks sit in each evidence tier.

    The numbers that say which assumptions are in play at all. A portfolio
    where every risk states its coverages needs neither an allocation nor a
    component split, and reporting that plainly is worth more than reporting
    either assumption carefully.
    """
    coverages = sum(1 for row in read.risks.rows if states_coverages(row))
    totals = sum(
        1
        for row in read.risks.rows
        if not states_coverages(row) and states_risk_total(row)
    )
    return {
        "risks": len(read.risks.rows),
        "coverages_stated": coverages,
        "risk_total_stated": totals,
        "allocated_from_policy": len(read.risks.rows) - coverages - totals,
    }


# -- canonical records -------------------------------------------------------------

#: Template column to the name the rest of the platform uses. The template is
#: written for a person and the engine is written for a model, and this is the
#: one place the two vocabularies meet: cohorts, allocation and promotion all
#: read canonical names, so none of them is shaped by what a spreadsheet column
#: happens to be called. A loading API that populates CASS directly produces
#: these same records without going near a workbook.
RISK_FIELDS: Mapping[str, str] = {
    "Policy ID": "business_id",
    "Risk reference": "location_number",
    "Risk name": "label",
    "Primary site": "primary_location",
    "Country": "country_code",
    "Latitude": "latitude",
    "Longitude": "longitude",
    "Address": "address",
    "Postal code": "postal_code",
    "Administrative area": "area_code",
    "Geocode precision": "precision",
    "Needs review": "needs_review",
    "Class of business": "class_of_business",
    "Total insured value": "location_tiv",
    "Occupancy": "occupancy_code",
    "Construction": "construction_code",
    "Year built": "year_built",
    "Storeys": "storeys",
    "Perils covered": "perils_covered",
    "Currency": "currency",
    "Building value": "BuildingTIV",
    "Other structures value": "OtherTIV",
    "Contents value": "ContentsTIV",
    "Business interruption value": "BITIV",
    "Risk deductible": "location_deductible",
    "Risk limit": "location_limit",
}

POLICY_FIELDS: Mapping[str, str] = {
    "Policy ID": "business_id",
    "Policy reference": "policy_id",
    "Currency": "currency",
    "Perils covered": "perils_covered",
    "Total insured value": "policy_tiv",
    "Inception date": "inception_date",
    "Expiry date": "expiry_date",
    "Layer": "layer_number",
    "Signed share": "signed_share",
    "Layer limit": "layer_limit",
    "Layer attachment": "layer_attachment",
    "Policy deductible": "policy_deductible",
    "Policy limit": "policy_limit",
}


CONTRACT_FIELDS: Mapping[str, str] = {
    "Contract number": "contract_number",
    "Layer": "layer_number",
    "Contract name": "name",
    "Contract type": "contract_type",
    "Inuring priority": "inuring_priority",
    "Ceded share": "ceded_percent",
    "Placed share": "placed_percent",
    "Attachment per event": "occurrence_attachment",
    "Limit per event": "occurrence_limit",
    "Risk level": "risk_level",
    "Limit per risk": "risk_limit",
    "Reinstatements": "reinstatements",
    "Reinstatement rate": "reinstatement_rate",
    "Reinstatement premium": "reinstatement_premium",
}

SCOPE_FIELDS: Mapping[str, str] = {
    "Contract number": "contract_number",
    "Policy ID": "business_id",
    "Policy reference": "policy_id",
    "Risk reference": "location_number",
    "Ceded share": "ceded_percent",
}


def risk_record(row: SourceRow) -> dict[str, Any]:
    """One risk row under the names the platform reads."""
    record = {name: row.values.get(column) for column, name in RISK_FIELDS.items()}
    record["row_number"] = row.row_number
    record["states_coverages"] = states_coverages(row)
    record["states_risk_total"] = states_risk_total(row)
    return record


def policy_record(row: SourceRow) -> dict[str, Any]:
    """One policy row under the names the platform reads."""
    record = {name: row.values.get(column) for column, name in POLICY_FIELDS.items()}
    record["row_number"] = row.row_number
    return record


def _record(row: SourceRow, fields: Mapping[str, str]) -> dict[str, Any]:
    record = {name: row.values.get(column) for column, name in fields.items()}
    record["row_number"] = row.row_number
    return record


def reinsurance_records(
    read: IntakeRead,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Both reinsurance sheets as canonical records, in sheet order."""
    return (
        [_record(row, CONTRACT_FIELDS) for row in read.contracts.rows],
        [_record(row, SCOPE_FIELDS) for row in read.scope.rows],
    )


def records(read: IntakeRead) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """The Risks and Policies sheets as canonical records, in sheet order."""
    return (
        [risk_record(row) for row in read.risks.rows],
        [policy_record(row) for row in read.policies.rows],
    )
