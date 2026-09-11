from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Daniel.Mhango\Documents\KRE_CATASTROPHE_MODEL")
OUTPUT_DIR = ROOT / "deliverables"
OUTPUT = OUTPUT_DIR / "Klapton Re Earthquake Catastrophe Modelling Platform Build Plan.docx"
LOGO = Path(r"C:\Users\Daniel.Mhango\Downloads\KRL PLC Logo-06 (1).png")

BLUE = "234A9E"
BLUE_PALE = "EAF0FA"
GREY = "808080"
GREY_DARK = "555555"
GREY_PALE = "F3F4F6"
LIGHT_BORDER = "D9D9D9"
WHITE = "FFFFFF"
BLACK = "000000"
FONT = "Arial"


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_font(run, name=FONT, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    set_font(run, size=8, color=GREY_DARK)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def no_split_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_paragraph(doc, text="", style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_font(lead, bold=True)
        body = p.add_run(text[len(bold_lead):])
        set_font(body)
    else:
        run = p.add_run(text)
        set_font(run)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        set_font(p.add_run(item))


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_font(p.add_run(item))


def add_table(doc, headers, rows, widths=None, font_size=9, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    no_split_row(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell, top=110, bottom=110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        set_font(run, size=font_size, bold=True, color=WHITE)
        if widths:
            set_cell_width(cell, widths[index])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        no_split_row(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else GREY_PALE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if index > 0 and len(str(value)) < 22:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_font(run, size=font_size, bold=(first_col_bold and index == 0))
            if widths:
                set_cell_width(cell, widths[index])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_section_title(doc, number, title, new_page=False):
    p = doc.add_paragraph(style="Heading 1")
    keep_with_next(p)
    run = p.add_run(f"{number}  {title}")
    set_font(run, size=18, bold=True, color=BLACK)
    return p


def add_subtitle(doc, title):
    p = doc.add_paragraph(style="Heading 2")
    keep_with_next(p)
    run = p.add_run(title)
    set_font(run, size=13, bold=True, color=BLACK)
    return p


def add_small_label(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    keep_with_next(p)
    run = p.add_run(text.upper())
    set_font(run, size=8.5, bold=True, color=BLUE)
    return p


def set_document_defaults(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size in (("Title", 28), ("Subtitle", 15), ("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if style_name.startswith("Heading") else 0)
        style.paragraph_format.space_after = Pt(5)
        p_pr = style._element.get_or_add_pPr()
        borders = p_pr.find(qn("w:pBdr"))
        if borders is not None:
            p_pr.remove(borders)

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)


def add_footer(section):
    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.94))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_cell_width(table.cell(0, 0), 6.1)
    set_cell_width(table.cell(0, 1), 0.84)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, top=40, bottom=20, start=20, end=20)
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    run = left.add_run("Klapton Reinsurance PLC  |  Earthquake Platform Build Plan")
    set_font(run, size=7.5, color=GREY_DARK)
    right = table.cell(0, 1).paragraphs[0]
    right.paragraph_format.space_after = Pt(0)
    add_page_number(right)
    table._tbl.remove(table._tbl.tblPr.first_child_found_in("w:tblBorders")) if table._tbl.tblPr.first_child_found_in("w:tblBorders") is not None else None


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    set_document_defaults(doc)

    # Cover page
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO), width=Inches(6.4))

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run("Klapton Re Earthquake Catastrophe Modelling Platform Build Plan"), name=FONT, size=28, bold=True, color=BLACK)

    p = doc.add_paragraph(style="Subtitle")
    p.paragraph_format.space_after = Pt(20)
    set_font(p.add_run("End to End Delivery Roadmap"), name=FONT, size=15, color=BLACK)

    metadata = [
        ("Prepared for", "Klapton Reinsurance PLC"),
        ("Planning date", "10 September 2026"),
        ("Document version", "1.0 Decision Baseline"),
        ("Initial peril", "Earthquake"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for label, value in metadata:
        row = table.add_row()
        set_cell_width(row.cells[0], 1.55)
        set_cell_width(row.cells[1], 4.9)
        for cell in row.cells:
            set_cell_margins(cell, top=65, bottom=65, start=0, end=80)
        p1 = row.cells[0].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_font(p1.add_run(label), size=9, bold=True, color=BLUE)
        p2 = row.cells[1].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        set_font(p2.add_run(value), size=9.5, color=GREY_DARK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    set_font(p.add_run("Purpose"), size=9, bold=True, color=BLUE)
    p = add_paragraph(doc, "This plan defines the product, scientific, data, engineering, validation and operating work required to build KRE's integrated catastrophe modelling platform. It establishes the target architecture and a staged route from the proven Docker test environment to a governed earthquake modelling service.")
    p.paragraph_format.space_after = Pt(0)

    doc.add_page_break()

    # Executive plan
    add_section_title(doc, "1", "Executive Direction")
    add_paragraph(doc, "KRE should build one browser-based catastrophe modelling product while preserving OpenQuake and Oasis as independently owned calculation engines. React will provide the analyst experience. Django will provide authentication, workflow orchestration, metadata, lineage and a stable KRE API. OpenQuake will generate earthquake hazard. A KRE-owned converter will transform supported OpenQuake outputs into versioned Oasis hazard packages. Oasis will apply vulnerability and financial structures and produce insured loss outputs.")
    add_paragraph(doc, "The platform will use Docker as the common execution boundary from development onward. Large scientific arrays will remain in HDF5, Parquet or Oasis binary artifacts stored outside the Django relational database. PostgreSQL will hold application records, references, checksums and audit history. This design avoids Windows path problems, prevents the user interface database from becoming a scientific array store, and allows each engine to be upgraded behind a tested adapter.")

    add_small_label(doc, "Build outcome")
    add_bullets(doc, [
        "An internal KRE application through which an analyst can select a model, upload and validate exposure, run earthquake hazard and loss analyses, monitor progress, inspect quality checks, compare results and export governed outputs.",
        "A reproducible earthquake model build process using stable geographic area-peril definitions, explicit event semantics, versioned vulnerability functions and documented financial assumptions.",
        "A portable container deployment that works on the current Windows workstation through Docker Desktop and can move to a managed Linux container environment without changing the product architecture.",
        "An extensible foundation on which additional countries and perils can be added through engine and model adapters rather than rewrites of the user interface.",
    ])

    add_small_label(doc, "Current evidence")
    add_paragraph(doc, "The official PiWind model has already completed end to end in the Oasis 2.5.7 model worker container against a Windows bind-mounted workspace. The test produced ground-up, insured and reinsurance outputs and passed the model validation check. This proves the selected local container boundary. It does not yet prove earthquake model correctness, production capacity or platform security; those are explicit workstreams in this plan.")

    add_small_label(doc, "Recommended release boundary")
    add_paragraph(doc, "Release one should be an internal KRE earthquake portfolio analysis platform. External client tenancy, self-service onboarding, billing, real-time event response and non-earthquake perils should be deferred until the internal workflow and model governance are proven.")

    add_section_title(doc, "2", "Product Scope", new_page=True)
    add_subtitle(doc, "Primary users")
    add_table(doc,
        ["User", "Primary need", "Release one capability"],
        [
            ["Catastrophe modeller", "Build and validate model versions", "Configure hazard, run conversion, register vulnerability, review scientific QA and publish model versions"],
            ["Portfolio analyst", "Run governed portfolio analyses", "Upload OED exposure, resolve validation issues, set analysis options, monitor jobs and review loss metrics"],
            ["Underwriter or pricing user", "Interpret decision metrics", "Review approved AAL, EP curves, event losses, geographic concentrations and comparison views"],
            ["Platform administrator", "Operate a reliable controlled service", "Manage users, roles, storage, engines, queues, retention and audit records"],
            ["Reviewer or approver", "Provide independent challenge", "Review assumptions, validation evidence, exceptions and model publication gates"],
        ],
        widths=[1.25, 2.0, 3.65],
        font_size=8.8,
        first_col_bold=True,
    )

    add_subtitle(doc, "Release one functional scope")
    add_bullets(doc, [
        "Authentication, role-based access and project workspaces.",
        "Model registry covering country, peril, engine version, hazard source, grid, vulnerability set, financial capability and publication state.",
        "OED location, account and reinsurance input upload with validation, error reports and controlled normalization.",
        "OpenQuake job creation, configuration validation, execution, progress, logs and artifact capture.",
        "OpenQuake to Oasis conversion with chunked HDF5 processing, deterministic identifiers, scientific QA and reproducible packaging.",
        "Oasis keys lookup, file generation and ground-up, insured and reinsurance loss calculation through the Oasis API.",
        "Results catalogue with average annual loss, exceedance probability curves, event loss tables, geographic summaries, run comparisons and downloadable OED or ORD-aligned outputs.",
        "Audit history, approvals, retention controls and operational dashboards.",
    ])

    add_subtitle(doc, "Deferred scope")
    add_bullets(doc, [
        "External multi-tenant client access and client-specific branding.",
        "Cyclone, flood and other perils, except for the extension points required to add them later.",
        "Full replacement of the native OpenQuake and Oasis administrative interfaces.",
        "Real-time post-event loss estimation and live sensor or agency feeds.",
        "Pricing workflow integration, policy administration integration and automated capital model submissions.",
        "Highly available multi-region deployment before workload and recovery requirements justify it.",
    ])

    add_section_title(doc, "3", "Product Experience", new_page=True)
    add_subtitle(doc, "Principal user journey")
    add_numbered(doc, [
        "Create or open a project and select an approved earthquake model version.",
        "Upload exposure and financial files directly to artifact storage through a secure upload session.",
        "Review OED validation, geocoding and model-coverage findings; correct or explicitly approve permitted exceptions.",
        "Choose analysis options from model-defined settings rather than editing engine configuration files.",
        "Submit the analysis and follow a single KRE status view while Django coordinates the OpenQuake, conversion and Oasis tasks.",
        "Review scientific and operational checks before results are released to decision users.",
        "Explore loss results, maps, curves and comparisons, then export an auditable result package.",
    ])

    add_subtitle(doc, "Core screens")
    add_table(doc,
        ["Screen", "Purpose", "Key content"],
        [
            ["Portfolio dashboard", "Show current work and exceptions", "Recent projects, run status, failed checks, storage and model notices"],
            ["Model catalogue", "Select an approved model", "Country, peril, version, publication state, assumptions and validation date"],
            ["Exposure workspace", "Prepare analysis inputs", "Upload status, OED errors, geocoding, taxonomy mapping, TIV summaries and unmapped records"],
            ["Analysis builder", "Configure a governed run", "Model settings, financial options, outputs, resource profile and validation summary"],
            ["Run monitor", "Explain progress and failure", "Pipeline stage, elapsed time, logs, warnings, artifacts and retry controls"],
            ["Results workspace", "Support risk interpretation", "AAL, EP curves, event tables, maps, uncertainty, comparisons and downloads"],
            ["Model build workspace", "Manage scientific assets", "Hazard runs, converter QA, vulnerability sets, model packages and approval gates"],
            ["Administration", "Operate the service", "Users, roles, engine health, queues, storage, retention and audit search"],
        ],
        widths=[1.3, 1.9, 3.7],
        font_size=8.6,
        first_col_bold=True,
    )

    add_subtitle(doc, "Visual design system")
    add_paragraph(doc, "The interface should use KRE blue #234A9E as the primary action, navigation and selection colour; white as the principal workspace surface; and a restrained neutral grey scale for borders, secondary text, inactive states and dense analytical views. Black remains the default reading colour. Charts should reserve KRE blue for the principal series and use accessible, colour-blind-safe secondary colours only where comparison requires them.")
    add_bullets(doc, [
        "Use a left navigation rail for stable product areas and a wide content canvas for maps, curves and tables.",
        "Make model version, portfolio, financial perspective and run state continuously visible.",
        "Translate engine terminology into analyst language while retaining expandable technical details for modellers.",
        "Use progressive disclosure for advanced OpenQuake and Oasis settings.",
        "Meet WCAG 2.1 AA contrast and keyboard-navigation expectations, including non-colour status cues.",
        "Treat long-running work as background jobs; never require the browser to remain open.",
    ])

    add_section_title(doc, "4", "Target Architecture", new_page=True)
    add_paragraph(doc, "The architecture separates the control plane from the scientific data plane. Django coordinates work and records what happened. Engine services perform calculations. Artifact storage carries large immutable inputs and outputs. Every boundary is accessed through a versioned adapter or supported API.")

    add_subtitle(doc, "End to end service flow")
    add_table(doc,
        ["KRE React", "KRE Django", "OpenQuake", "KRE Converter", "Oasis"],
        [[
            "Analyst workflow and results",
            "Identity, orchestration, metadata and audit",
            "Earthquake hazard and supported exports",
            "HDF5 to events, occurrence and footprint",
            "Vulnerability, financial terms and loss outputs",
        ]],
        widths=[1.32, 1.45, 1.3, 1.45, 1.38],
        font_size=8.2,
    )
    add_paragraph(doc, "Shared infrastructure: PostgreSQL for KRE application records; an object store for uploads, HDF5, model packages and results; a background task queue; centralized logs and metrics; and a secrets service. The OpenQuake and Oasis internal stores remain engine-owned.")

    add_subtitle(doc, "Component responsibilities")
    add_table(doc,
        ["Component", "Responsibility", "Boundary rule"],
        [
            ["React TypeScript application", "Analyst workflow, validation presentation, mapping, charting and result comparison", "Calls only the KRE API; does not call engine APIs directly"],
            ["Django REST API", "Users, roles, projects, runs, model registry, metadata, approvals and audit", "Stores references to large artifacts, not scientific arrays"],
            ["Workflow workers", "Execute durable jobs, retries, cancellation and stage transitions", "Every task is idempotent and records input and output checksums"],
            ["OpenQuake adapter", "Submit, monitor and export calculations through supported interfaces", "No KRE writes to OpenQuake internal tables"],
            ["Converter service", "Translate versioned OpenQuake output into validated Oasis model data", "Pinned compatibility matrix and deterministic output"],
            ["Oasis adapter", "Register models, upload exposure, run analyses and collect outputs", "Use the Platform API rather than modifying the Oasis UI"],
            ["Artifact store", "Hold large immutable and versioned files", "Content checksum, retention class and access policy on every object"],
            ["PostgreSQL", "Hold KRE control-plane records", "No event-site-IMT observation table"],
            ["Observability stack", "Aggregate structured logs, metrics, traces and alerts", "Correlation ID follows every run across all services"],
        ],
        widths=[1.45, 3.15, 2.3],
        font_size=8.3,
        first_col_bold=True,
    )

    add_subtitle(doc, "Recommended technology baseline")
    add_bullets(doc, [
        "React with TypeScript and a documented KRE component library; MapLibre GL JS for maps; a mature charting library for EP and loss distributions.",
        "Django, Django REST Framework and generated OpenAPI contracts; PostgreSQL for the KRE database.",
        "Celery-compatible durable background execution. Redis is acceptable for development; RabbitMQ or another production-grade broker should be selected through the deployment decision.",
        "S3-compatible object storage. A local filesystem-backed service is acceptable for the first workstation deployment if the same object interface is preserved.",
        "Docker Compose for local development and the first controlled deployment; managed containers or Kubernetes only when availability and workload justify the operating overhead.",
        "Pinned container image digests, software bills of materials, vulnerability scanning and a tested engine compatibility matrix.",
    ])

    add_section_title(doc, "5", "Application and Data Model", new_page=True)
    add_subtitle(doc, "KRE control-plane records")
    add_table(doc,
        ["Record", "Purpose", "Examples of retained fields"],
        [
            ["Project", "Business workspace", "Owner, team, purpose, access, status"],
            ["Exposure version", "Immutable input version", "OED version, file references, schema result, TIV summaries, checksum"],
            ["Model version", "Published calculation capability", "Country, peril, engine versions, grid, hazard, vulnerability and approval state"],
            ["Hazard run", "OpenQuake execution record", "Calculation ID, settings hash, image digest, state, timings and artifact references"],
            ["Conversion run", "OpenQuake to Oasis lineage", "Converter version, event policy, binning, source and target checksums, QA state"],
            ["Analysis run", "Portfolio loss execution", "Exposure, model, settings, financial perspective, queue state and Oasis identifiers"],
            ["Result set", "Discoverable approved outputs", "ORD artifacts, AAL, validation state, publication and retention"],
            ["Audit event", "Governance evidence", "Actor, action, timestamp, object, before and after references"],
        ],
        widths=[1.35, 2.1, 3.45],
        font_size=8.5,
        first_col_bold=True,
    )

    add_subtitle(doc, "Artifact lifecycle")
    add_table(doc,
        ["Artifact class", "Default retention", "Treatment"],
        [
            ["Source inputs and manifests", "Permanent for published model versions", "Immutable, checksummed and access controlled"],
            ["OpenQuake HDF5 GMF", "Until footprint acceptance, then policy-based archive or expiry", "Chunked binary format; never expanded into Django rows"],
            ["Diagnostic CSV", "Short-lived", "Created only for inspection or interoperability exceptions"],
            ["Accepted Oasis model package", "Versioned model asset", "Immutable release with scientific approval and compatibility record"],
            ["Portfolio input versions", "Business retention policy", "Encrypted, access controlled and deletable without damaging model assets"],
            ["Loss result package", "Business and regulatory policy", "ORD-aligned outputs plus KRE summaries and lineage"],
            ["Logs and operational metrics", "Tiered by usefulness", "Short hot retention, longer error and audit retention"],
        ],
        widths=[1.55, 2.25, 3.1],
        font_size=8.5,
        first_col_bold=True,
    )

    add_subtitle(doc, "File transfer principles")
    add_bullets(doc, [
        "The browser uploads large files directly to object storage with a short-lived signed upload session.",
        "Django records the artifact only after checksum and malware or content validation completes.",
        "Workers read and write through the artifact interface; Windows host paths do not appear in calculation contracts.",
        "HDF5 and Parquet are used for large typed arrays; CSV remains a human-inspection and interchange fallback.",
        "A calculation can be reconstructed from retained inputs, versions, settings and checksums even when an intermediate GMF has expired.",
    ])

    add_section_title(doc, "6", "Earthquake Model Engineering", new_page=True)
    add_subtitle(doc, "Hazard model basis")
    add_paragraph(doc, "The first production model should use a stable, versioned spatial mesh that is independent of any uploaded portfolio. The mesh may be a regular grid, variable-resolution grid or a defensible polygon scheme. Each cell receives a permanent integer area-peril identifier. Exposure is mapped to this mesh by the Oasis keys service. This allows one accepted hazard package to serve many portfolio analyses.")

    add_small_label(doc, "Country model definition")
    add_bullets(doc, [
        "Document the source model, licence, tectonic region logic, source-model branches and minimum magnitude.",
        "Select ground-motion models, logic-tree weights, site parameters, correlation assumptions and truncation settings with scientific justification.",
        "Define the stable site mesh and any interpolation or nearest-cell policy, including treatment near national borders and coastlines.",
        "Select intensity measures needed by the vulnerability functions. Avoid producing unused IMTs merely because OpenQuake can calculate them.",
        "Define investigation time, stochastic event-set count, random seed policy and occurrence representation.",
        "Establish hazard benchmarks against published maps, curves or independent OpenQuake calculations before converter development is accepted.",
    ])

    add_subtitle(doc, "Area-peril grid design spike")
    add_table(doc,
        ["Decision factor", "Assessment"],
        [
            ["Resolution", "Balance local hazard gradients and geocoding uncertainty against event-site storage and runtime"],
            ["Site conditions", "Determine whether Vs30 and other parameters are cell attributes, exposure attributes or scenario options"],
            ["Variable resolution", "Consider finer cells in high-exposure or high-gradient areas and coarser cells elsewhere"],
            ["Version stability", "Never silently change an existing area-peril ID; publish a new grid version"],
            ["Mapping QA", "Measure unmapped, boundary, offshore, low-confidence and duplicate-location cases"],
            ["Reuse", "The grid should support model-building and portfolio-running without regenerating hazard for each portfolio"],
        ],
        widths=[1.7, 5.2],
        font_size=8.8,
        first_col_bold=True,
    )

    add_subtitle(doc, "Event semantics")
    add_paragraph(doc, "Before code is written, KRE must decide how an OpenQuake rupture, stochastic occurrence, realization and ground-motion sample map to an Oasis event and its occurrence record. A defensible design may represent each simulated occurrence as a separate Oasis event, or aggregate repeated ground-motion samples into intensity-bin probabilities for a rupture-level event. The selected design must preserve annual frequency, uncertainty and correlation in a way that is consistent with the intended Oasis sampling workflow.")
    add_bullets(doc, [
        "Publish a formal event identity specification and worked examples.",
        "Prove that total occurrence rates and effective time are preserved after conversion.",
        "Prove that correlated ground motion is not accidentally converted into independent site sampling.",
        "Record event and realization lineage so any material loss event can be traced back to OpenQuake output.",
    ])

    add_section_title(doc, "7", "OpenQuake to Oasis Converter", new_page=True)
    add_paragraph(doc, "The converter is a KRE product component, not an ad hoc export script. It should be packaged as a stateless container and called by a durable background job. Its public contract is a conversion manifest; its output is a complete candidate Oasis model package plus machine-readable validation evidence.")

    add_subtitle(doc, "Converter inputs and outputs")
    add_table(doc,
        ["Input", "Transformation", "Output"],
        [
            ["Supported OpenQuake HDF5 GMF export", "Chunk by event and site without loading the complete dataset", "Footprint records grouped by event and area peril"],
            ["OpenQuake event and rupture metadata", "Apply the approved event identity and occurrence policy", "events and occurrence data with preserved frequency"],
            ["Stable site mesh", "Map site identifiers to permanent area-peril identifiers", "Area-peril dictionary and lookup assets"],
            ["Intensity specification", "Convert IMT values into versioned intensity bins", "Footprint intensity-bin probabilities"],
            ["Conversion manifest", "Validate versions, schemas, counts and checksums", "Signed or approved model-build evidence"],
        ],
        widths=[2.0, 2.75, 2.15],
        font_size=8.5,
    )

    add_subtitle(doc, "Engineering requirements")
    add_bullets(doc, [
        "Use the supported OpenQuake export boundary. Direct native-datastore access is permitted only through a version-specific adapter with compatibility tests.",
        "Process data in configurable chunks and expose memory, throughput and spill-to-disk metrics.",
        "Generate deterministic identifiers and byte-stable outputs when inputs, versions and settings are unchanged.",
        "Validate schema, event counts, area-peril counts, IMTs, intensity ranges, probabilities, occurrence rates, duplicate keys and orphan records.",
        "Write a conversion report containing checksums, warnings, excluded records, runtime and peak memory.",
        "Support restart from durable checkpoints for long conversions without producing partially published model packages.",
        "Keep mapping, binning and event policies configurable through versioned files rather than source-code edits.",
    ])

    add_subtitle(doc, "Scientific acceptance tests")
    add_bullets(doc, [
        "Reconstruct selected OpenQuake hazard distributions from the candidate footprint and compare them within approved tolerances.",
        "Compare hazard curves at a representative set of high, medium and low hazard cells.",
        "Check annual event frequency and period weighting before and after conversion.",
        "Test spatial coherence for selected events and quantify effects introduced by grid and intensity discretisation.",
        "Run controlled vulnerability functions to isolate hazard conversion from financial-model effects.",
        "Maintain small golden fixtures plus a country-scale regression dataset.",
    ])

    add_section_title(doc, "8", "Oasis Model and Loss Workflow", new_page=True)
    add_subtitle(doc, "Model package")
    add_paragraph(doc, "The earthquake model release combines the converted hazard module, a versioned vulnerability module, occurrence definitions, lookup configuration, model settings and supporting dictionaries. Oasis relates events and area perils through the footprint and relates intensities to damage through vulnerability functions. KRE should publish these together as a single governed release even when individual components have separate scientific owners.")

    add_subtitle(doc, "Vulnerability workstream")
    add_bullets(doc, [
        "Define the KRE exposure taxonomy and an explicit mapping from OED occupancy, construction, height, age and other attributes.",
        "Acquire or develop vulnerability functions with documented provenance and permitted commercial use.",
        "Align every vulnerability set with the converter's intensity measure and bin definitions.",
        "Represent damage-ratio uncertainty deliberately and test sensitivity to vulnerability selection.",
        "Create fallback and unknown-taxonomy policies that are visible to analysts and included in result caveats.",
        "Approve vulnerability versions independently from hazard versions, then publish tested combinations in the model registry.",
    ])

    add_subtitle(doc, "Exposure and financial workflow")
    add_bullets(doc, [
        "Use current OED as the external input contract and validate files before Oasis file generation.",
        "Implement a model-specific keys service that returns area-peril and vulnerability identifiers plus precise failure reasons.",
        "Support ground-up loss first, followed by insured loss and then reinsurance once relevant financial structures have golden tests.",
        "Expose supported Oasis analysis settings through model-defined UI controls and safe defaults.",
        "Prefer current ORD-aligned result outputs for downstream analytics and exports.",
        "Preserve the native Oasis result package alongside KRE summaries so advanced users can independently inspect outputs.",
    ])

    add_subtitle(doc, "Analysis perspectives")
    add_table(doc,
        ["Perspective", "Release condition", "Primary outputs"],
        [
            ["Ground-up loss", "Hazard, vulnerability and exposure mapping approved", "AAL, EP curves, event loss and geographic summaries"],
            ["Insured loss", "Supported OED policy terms pass financial golden tests", "Insured AAL, EP curves, policy and portfolio summaries"],
            ["Reinsurance loss", "Treaty scope and inuring logic pass representative contract tests", "Ceded and net metrics, treaty summaries and recoveries"],
        ],
        widths=[1.4, 3.15, 2.35],
        font_size=8.8,
        first_col_bold=True,
    )

    add_section_title(doc, "9", "Security Governance and Licensing", new_page=True)
    add_subtitle(doc, "Security controls")
    add_bullets(doc, [
        "Use centralized identity, multi-factor authentication where available, least-privilege roles and short-lived service credentials.",
        "Encrypt network traffic and stored portfolio artifacts; isolate engine networks from direct user access.",
        "Scan uploads, container images and dependencies; produce a software bill of materials for releases.",
        "Keep secrets outside source control and container images; rotate them through the deployment platform.",
        "Record access, download, deletion, model publication, exception approval and configuration changes in an immutable audit trail.",
        "Apply per-project authorization to both metadata and artifact retrieval; a guessed object key must never grant access.",
        "Define incident response, backup, restore and controlled deletion procedures before production use.",
    ])

    add_subtitle(doc, "Model governance")
    add_table(doc,
        ["Gate", "Required evidence", "Approver"],
        [
            ["Hazard candidate", "Source, GMM, logic tree, site model, grid and benchmark report", "Earthquake hazard reviewer"],
            ["Converter candidate", "Compatibility, lineage, frequency, discretisation and regression evidence", "Model engineering reviewer"],
            ["Vulnerability candidate", "Provenance, taxonomy mapping, uncertainty and sensitivity evidence", "Vulnerability reviewer"],
            ["Model release", "End-to-end loss validation, performance, known limitations and reproducible package", "KRE model owner"],
            ["Platform release", "Security, restore, monitoring, UAT and rollback evidence", "Product and technology owners"],
        ],
        widths=[1.45, 4.1, 1.35],
        font_size=8.4,
        first_col_bold=True,
    )

    add_subtitle(doc, "Licensing")
    add_paragraph(doc, "OpenQuake is distributed under the GNU Affero General Public License version 3, while the main OasisLMF repository uses a BSD licence. KRE should obtain legal review before offering the platform to external network users, distributing modified engine containers or embedding third-party hazard and vulnerability data. The design should prefer unmodified upstream engine containers and separate KRE adapters, but service separation must not be treated as a substitute for licence analysis.")

    add_section_title(doc, "10", "Delivery and Operations", new_page=True)
    add_subtitle(doc, "Environment progression")
    add_table(doc,
        ["Environment", "Purpose", "Minimum characteristics"],
        [
            ["Developer", "Feature development and small fixtures", "Docker Compose, seeded data, isolated buckets and local observability"],
            ["Integration", "Cross-service and engine compatibility", "Production-like images, automated PiWind and earthquake fixtures, disposable data"],
            ["Validation", "Scientific model acceptance", "Controlled model assets, representative compute, reviewer access and retained evidence"],
            ["User acceptance", "Business workflow approval", "Masked or approved data, role testing, realistic portfolios and support procedures"],
            ["Production", "Governed analyses", "Backups, monitoring, secret management, capacity controls, restore tests and change approval"],
        ],
        widths=[1.2, 2.3, 3.4],
        font_size=8.6,
        first_col_bold=True,
    )

    add_subtitle(doc, "Capacity approach")
    add_paragraph(doc, "The PiWind run used approximately 21.6 GB at peak under a highly parallel all-output configuration. The production design must expose compute profiles rather than inheriting maximum parallelism. The earthquake pilot should measure OpenQuake generation, conversion and Oasis loss workloads separately and use those measurements to set queue concurrency, worker memory, storage tiers and run-size limits.")
    add_bullets(doc, [
        "Define small, standard and large execution profiles with explicit CPU, memory and timeout limits.",
        "Prevent multiple large jobs from exhausting the workstation or production node.",
        "Instrument event throughput, GMF bytes, conversion throughput, Oasis event throughput and output volume.",
        "Allow cancellation and clean restart without leaving published partial results.",
        "Set storage quotas and lifecycle rules by artifact class rather than deleting files manually.",
    ])

    add_subtitle(doc, "Operational service levels")
    add_paragraph(doc, "During the internal pilot, the service level should prioritize recoverability and transparency over continuous availability. Target values for availability, recovery time, recovery point, support hours and maximum run duration should be agreed after pilot workload measurements. Production approval requires a successful restore exercise, not merely confirmation that backups exist.")

    add_section_title(doc, "11", "Quality and Validation Strategy", new_page=True)
    add_table(doc,
        ["Test layer", "Purpose", "Representative evidence"],
        [
            ["Unit", "Verify local transformation and business rules", "Binning boundaries, IDs, status transitions, permissions and financial mappings"],
            ["Contract", "Detect upstream API and schema change", "Pinned OpenQuake and Oasis request, response and export fixtures"],
            ["Component", "Verify each container with real dependencies", "HDF5 chunking, keys lookup, model packaging and result ingestion"],
            ["End to end", "Prove the complete analyst workflow", "PiWind baseline and a small earthquake golden model through the KRE UI and APIs"],
            ["Scientific", "Establish model credibility", "Hazard benchmarks, event frequency, spatial checks, vulnerability sensitivity and loss comparisons"],
            ["Performance", "Set safe limits and capacity", "Large GMF, portfolio and result tests with CPU, memory, I/O and duration"],
            ["Security", "Protect portfolio and model assets", "Authorization, upload, secrets, dependency, image and penetration testing"],
            ["Resilience", "Prove recovery behavior", "Worker termination, engine failure, object-store interruption, retry and restore exercises"],
        ],
        widths=[1.15, 2.2, 3.55],
        font_size=8.4,
        first_col_bold=True,
    )

    add_subtitle(doc, "Release acceptance")
    add_bullets(doc, [
        "A fresh environment can be built from version-controlled configuration and pinned images.",
        "An approved user can complete the full journey without editing engine files or using a command line.",
        "Every result is traceable to immutable exposure, model, engine, converter and settings versions.",
        "Engine or worker failure produces an intelligible state and a safe retry or cancellation path.",
        "Scientific benchmarks and converter tolerances are approved by an independent reviewer.",
        "Ground-up and supported financial outputs match golden results within documented tolerances.",
        "Access control, audit, backup and restore tests pass.",
        "Known limitations are visible in the model catalogue and exported result package.",
    ])

    add_section_title(doc, "12", "Phased Delivery Roadmap", new_page=True)
    add_paragraph(doc, "The roadmap is organized around working vertical slices. Durations are planning ranges for a core team of approximately five to seven people with dedicated catastrophe modelling input. Scientific data acquisition, licensing or independent validation can extend the schedule and should be tracked separately from software delivery.")
    add_table(doc,
        ["Phase", "Indicative duration", "Outcome and exit gate"],
        [
            ["0 Direction and model charter", "2 to 3 weeks", "Confirm release users, pilot country, grid policy, event semantics study, data rights, deployment target and acceptance authority"],
            ["1 Platform foundation", "3 to 5 weeks", "React and Django skeleton, identity, projects, PostgreSQL, object storage, queue, observability and automated environments"],
            ["2 Engine vertical slice", "4 to 6 weeks", "KRE submits and monitors OpenQuake and Oasis jobs; PiWind runs from the KRE interface; artifacts and lineage are visible"],
            ["3 Earthquake hazard prototype", "6 to 10 weeks", "Pilot source model, stable grid, OpenQuake settings, benchmark calculation and capacity measurements approved"],
            ["4 Converter and model package", "8 to 12 weeks", "Chunked HDF5 conversion, event and occurrence mapping, footprint generation, QA report and golden package approved"],
            ["5 Vulnerability and loss workflow", "6 to 10 weeks", "Taxonomy mapping, vulnerability set, keys service, ground-up loss and supported financial calculations validated"],
            ["6 Analyst product", "6 to 9 weeks", "Exposure workspace, analysis builder, run monitor, result views, maps, comparisons, exports and model catalogue complete"],
            ["7 Production readiness", "5 to 8 weeks", "Security, performance, restore, UAT, operating procedures, training, limitations and release approvals complete"],
            ["8 Controlled pilot", "4 to 6 weeks", "Selected KRE portfolios run under supervision; findings resolved and operating thresholds confirmed"],
        ],
        widths=[1.45, 1.35, 4.1],
        font_size=8.2,
        first_col_bold=True,
    )
    add_paragraph(doc, "With overlap between platform, interface and model work, an internal earthquake MVP is likely to require approximately seven to ten months after the model charter is approved. A fully governed production release is more realistically a nine to twelve month programme. These are planning ranges, not delivery commitments.")

    add_subtitle(doc, "Milestone demonstrations")
    add_table(doc,
        ["Milestone", "Demonstration"],
        [
            ["M1 Foundation", "Sign in, create a project, upload an artifact and observe a durable background task"],
            ["M2 Engine integration", "Run PiWind end to end from KRE and inspect lineage and outputs"],
            ["M3 Hazard", "Run the pilot OpenQuake model on the stable grid and compare benchmark hazard"],
            ["M4 Conversion", "Produce and validate an Oasis footprint without CSV staging or database array ingestion"],
            ["M5 Loss", "Run one approved portfolio through ground-up and insured loss"],
            ["M6 Product", "Complete the analyst journey and compare two governed runs"],
            ["M7 Production", "Restore the platform and a selected run from backup in a clean environment"],
        ],
        widths=[1.5, 5.4],
        font_size=8.8,
        first_col_bold=True,
    )

    add_section_title(doc, "13", "Team and Delivery Governance", new_page=True)
    add_table(doc,
        ["Role", "Core accountability", "Indicative involvement"],
        [
            ["KRE product owner", "Scope, priorities, business acceptance and stakeholder decisions", "Dedicated"],
            ["Catastrophe model owner", "Scientific requirements, assumptions, validation and model approval", "Dedicated during model phases"],
            ["Technical lead", "Architecture, contracts, code quality, security and release design", "Dedicated"],
            ["Backend engineer", "Django API, workflow, engine adapters, lineage and operations", "One to two dedicated"],
            ["Frontend engineer", "React workflow, maps, charts, accessibility and design system", "One dedicated"],
            ["Scientific or data engineer", "HDF5 conversion, geospatial grid, model packaging and performance", "One dedicated"],
            ["QA and automation engineer", "Test strategy, fixtures, performance, resilience and UAT evidence", "Dedicated from foundation onward"],
            ["Platform engineer", "Containers, CI, environments, observability, backup and security", "Part-time then dedicated near release"],
            ["Independent reviewers", "Hazard, vulnerability, financial and security challenge", "At defined gates"],
        ],
        widths=[1.45, 3.8, 1.65],
        font_size=8.4,
        first_col_bold=True,
    )

    add_subtitle(doc, "Ways of working")
    add_bullets(doc, [
        "Maintain one prioritized product backlog but separate software completion from scientific approval.",
        "Record architecture, model and data decisions in short version-controlled decision records.",
        "Demonstrate a working vertical slice at least every two weeks.",
        "Require review for engine-version changes, event semantics, grid versions, vulnerability changes and financial interpretation.",
        "Treat model packages as releases with changelogs, compatibility, evidence and rollback—not as folders copied between machines.",
        "Keep upstream contributions separate from KRE product delivery; propose general fixes upstream where practical.",
    ])

    add_section_title(doc, "14", "Risk Register", new_page=True)
    add_table(doc,
        ["Risk", "Impact", "Mitigation and trigger"],
        [
            ["Event semantics are defined too late", "Converter rework and invalid annual frequency", "Complete and approve the event identity study before production converter development"],
            ["Hazard grid is portfolio-specific", "Repeated GMF generation, high storage and inconsistent comparisons", "Adopt a stable area-peril grid and version it before model build"],
            ["Third-party data rights are unclear", "Model cannot be deployed or shared", "Complete a source and vulnerability licence register in phase zero"],
            ["OpenQuake or Oasis upgrade breaks integration", "Interrupted releases or changed results", "Pin image digests, maintain contract tests and upgrade through a compatibility environment"],
            ["Large files pass through Django or CSV", "Poor performance, memory pressure and fragile workflows", "Direct object uploads and chunked HDF5 or Parquet processing"],
            ["Scientific validation is treated as software QA", "Technically correct but unreliable model", "Separate approval gates and independent reviewers"],
            ["Unbounded concurrency exhausts the host", "Failed jobs and poor user experience", "Resource profiles, queue limits, admission control and measured capacity"],
            ["Custom engine forks accumulate", "Expensive maintenance and delayed security upgrades", "Use supported APIs and isolate KRE logic in adapters and converter services"],
            ["Financial terms are assumed rather than tested", "Material insured or ceded loss errors", "Representative OED contract fixtures and actuarial review"],
            ["Result interface hides uncertainty", "Overconfident decisions", "Display model version, uncertainty, data quality and limitations beside decision metrics"],
        ],
        widths=[2.0, 1.8, 3.1],
        font_size=8.1,
        first_col_bold=True,
    )

    add_section_title(doc, "15", "Decisions Required", new_page=True)
    add_paragraph(doc, "The following decisions should be completed during the model charter. Defaults shown here allow engineering to begin but should not be treated as silent assumptions.")
    add_table(doc,
        ["Decision", "Recommended default", "Why it matters"],
        [
            ["Release users", "KRE staff only", "Determines tenancy, authentication, licensing and support obligations"],
            ["Pilot country", "Select one country with accessible source, site and validation data", "Controls scientific workload and representative portfolios"],
            ["Hazard spatial basis", "Stable country grid or variable-resolution mesh", "Determines GMF reuse, storage, keys lookup and accuracy"],
            ["Event representation", "Formal study before converter build", "Controls frequency, uncertainty, correlation and footprint probabilities"],
            ["Production hosting", "Portable containers; decide host after capacity spike", "Controls broker, object store, secrets, backup and identity integration"],
            ["Portfolio scale", "Measure small, median and largest expected portfolios", "Sets upload, mapping, queue, storage and result-view requirements"],
            ["Vulnerability source", "Use only functions with documented provenance and commercial rights", "Controls model credibility, cost and publication rights"],
            ["Financial scope", "Ground-up first; add insured then reinsurance behind golden tests", "Avoids mixing hazard validation with complex contract interpretation"],
            ["Retention", "Keep manifests and published packages; expire raw GMFs after acceptance unless governed otherwise", "Balances auditability, reproducibility and storage cost"],
            ["Availability target", "Recoverable business-hours service for internal pilot", "Avoids premature high-availability complexity"],
        ],
        widths=[1.65, 2.45, 2.8],
        font_size=8.2,
        first_col_bold=True,
    )

    add_section_title(doc, "16", "First Thirty Days", new_page=True)
    add_numbered(doc, [
        "Approve the internal release boundary, pilot country selection criteria and model approval roles.",
        "Create the version-controlled repository structure for the KRE web application, engine adapters, converter, deployment configuration, model specifications and test fixtures.",
        "Build the Docker Compose foundation with Django, React, PostgreSQL, artifact storage, queue, OpenQuake and Oasis services.",
        "Turn the successful PiWind command-line test into an automated integration test and trigger it through a minimal KRE workflow.",
        "Run an OpenQuake API and HDF5 export spike using a small event-based earthquake calculation.",
        "Produce the event identity decision paper and a first area-peril grid experiment for the proposed pilot country.",
        "Define the KRE run state machine, artifact manifest, engine adapter interfaces and compatibility matrix.",
        "Create the first clickable React workflow for project, exposure upload, analysis submission, run monitoring and result navigation.",
        "Establish CI checks for unit tests, contracts, container builds, security scans and the PiWind baseline.",
        "Agree the evidence required to accept the first earthquake vertical slice.",
    ])

    add_subtitle(doc, "Thirty day exit evidence")
    add_bullets(doc, [
        "A clean-machine setup can launch the development platform from documented commands.",
        "KRE can submit one Oasis smoke analysis and one small OpenQuake calculation through service adapters.",
        "All produced artifacts have a recorded URI, checksum, owner, status and retention class.",
        "The fixed-grid and event-semantics decisions have named owners and dated review points.",
        "The programme backlog and milestone demonstrations are estimated by the delivery team.",
    ])

    add_section_title(doc, "17", "Reference Baseline", new_page=True)
    add_paragraph(doc, "The following upstream sources define the initial implementation boundary. KRE should pin exact releases and container digests after compatibility testing rather than following moving latest tags.")

    refs = [
        ("OpenQuake Engine repository and release baseline", "https://github.com/gem/oq-engine", "At the planning date, the project identifies 3.23 as the long-term-support line and 3.26 as the latest stable line."),
        ("OpenQuake engine architecture", "https://docs.openquake.org/oq-engine/master/manual/contributing/architecture.html", "Describes per-calculation HDF5 datastores and the separate metadata database."),
        ("OpenQuake REST API", "https://docs.openquake.org/oq-engine/manual/latest/api-reference/rest-api.html", "Provides the supported remote calculation and result boundary."),
        ("OpenQuake export guidance", "https://docs.openquake.org/oq-engine/manual/master/user-guide/advanced/useful-oq-commands.html", "Lists GMF HDF5 export and recommends binary formats for large outputs."),
        ("OasisLMF repository", "https://github.com/OasisLMF/OasisLMF", "Defines the supported 2.5.x line for 2026 and the model development toolkit."),
        ("Oasis Platform API reference", "https://oasislmf.github.io/2.5.7/platform/reference/index.html", "Documents the Platform REST APIs and OpenAPI schemas."),
        ("Oasis model data formats", "https://oasislmf.github.io/2.5.7/oasislmf/reference/Oasis-model-data-formats.html", "Defines events, area perils, footprints, intensity, vulnerability and occurrence concepts."),
        ("Oasis keys service", "https://oasislmf.github.io/2.5.7/oasislmf/explanation/keys-service.html", "Defines the exposure-to-model lookup boundary."),
        ("Oasis result formats", "https://oasislmf.github.io/2.5.7/oasislmf/reference/outputs/results.html", "Documents current ORD-aligned loss result outputs."),
    ]
    for title, url, description in refs:
        p = doc.add_paragraph(style="List Bullet")
        add_hyperlink(p, title, url)
        run = p.add_run(f". {description}")
        set_font(run, size=9.5)

    add_subtitle(doc, "Version policy")
    add_paragraph(doc, "The production selection should favour the OpenQuake LTS line unless a required model or export capability is available only in the newer stable line. Oasis should use a tested 2.5.x patch release and matching Platform and model-worker images. Every engine update must run contract, scientific regression and PiWind or earthquake end-to-end suites before promotion.")

    # Footer applies to all sections; cover also receives discreet footer.
    for section in doc.sections:
        add_footer(section)

    core = doc.core_properties
    core.title = "Klapton Re Earthquake Catastrophe Modelling Platform Build Plan"
    core.subject = "End to end product, model, engineering and delivery plan"
    core.author = "Klapton Reinsurance PLC"
    core.keywords = "Klapton Re, catastrophe modelling, earthquake, OpenQuake, OasisLMF, React, Django"
    core.comments = ""

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
